#!/usr/bin/env python3
"""Manage approved per-client UTM workbooks and rebuild app dictionaries."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urlsplit

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as error:  # pragma: no cover - environment guidance
    raise SystemExit("openpyxl is required. Install it with: python -m pip install openpyxl") from error


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
DEFAULT_CLIENTS_DIR = SCRIPT_DIR / "clients"
APPROVED_SHEET = "Approved UTMs"
WORKBOOK_COLUMNS = (
    "client", "destination_url", "source", "medium", "campaign",
    "term", "content", "bitly", "creation_date",
)
REQUIRED_COLUMNS = ("client", "destination_url", "source", "medium", "campaign")
UTM_FIELDS = ("campaign", "source", "medium", "term", "content")
VALUE_COUNT_COLUMNS = ("field", "campaign", "count", "unique_clients", "source", "medium", "term", "content")
MASTER_COLUMNS = (
    "creation_date", "client_code", "destination_url", "source", "medium",
    "campaign", "term", "content", "hashtag", "bitly", "source_file",
    "sheet_name", "excel_row", "canonical_key",
)
INPUT_ALIASES = {
    "client": ("client", "client_code"),
    "destination_url": ("destination_url",),
    "source": ("utm_source", "source"),
    "medium": ("utm_medium", "medium"),
    "campaign": ("utm_campaign", "campaign", "canonical_campaign"),
    "term": ("utm_term", "term"),
    "content": ("utm_content", "content"),
    "bitly": ("short_url", "bitly"),
    "creation_date": ("first_seen_at", "creation_date"),
}


class ValidationError(ValueError):
    """Raised when an approved input is unsafe to publish."""


@dataclass(frozen=True)
class ApprovedRow:
    client: str
    destination_url: str
    source: str
    medium: str
    campaign: str
    term: str = ""
    content: str = ""
    bitly: str = ""
    creation_date: str = ""
    source_file: str = ""
    excel_row: int = 0

    def normalized(self, value: str) -> str:
        return clean(value).lower()

    @property
    def key(self) -> tuple[str, ...]:
        return tuple(self.normalized(value) for value in (
            self.client, self.destination_url, self.source, self.medium,
            self.campaign, self.term, self.content,
        ))

    def master_dict(self) -> dict[str, str | int]:
        client, destination, source, medium, campaign, term, content = self.key
        canonical_key = "|".join((destination, source, medium, campaign, term, content, client))
        return {
            "creation_date": self.creation_date,
            "client_code": client,
            "destination_url": destination,
            "source": source,
            "medium": medium,
            "campaign": campaign,
            "term": term,
            "content": content,
            "hashtag": "",
            "bitly": clean(self.bitly),
            "source_file": self.source_file,
            "sheet_name": APPROVED_SHEET,
            "excel_row": self.excel_row,
            "canonical_key": canonical_key,
        }


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def valid_http_url(value: str) -> bool:
    try:
        parsed = urlsplit(value)
        return parsed.scheme.lower() in {"http", "https"} and bool(parsed.netloc)
    except ValueError:
        return False


def normalized_date(value: object) -> str:
    text = clean(value)
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError as error:
        raise ValidationError(f"invalid creation_date {text!r}; use an ISO date") from error


def input_value(row: dict[str, str], field: str) -> str:
    for alias in INPUT_ALIASES[field]:
        value = clean(row.get(alias))
        if value:
            return value
    return ""


def import_csv(input_path: Path, client: str, clients_dir: Path) -> Path:
    client = clean(client)
    if not client or not re.fullmatch(r"[A-Za-z0-9_-]+", client):
        raise ValidationError("client must contain only letters, numbers, underscores, or hyphens")
    with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValidationError(f"CSV has no header: {input_path}")
        rows = []
        for number, source_row in enumerate(reader, start=2):
            source_client = input_value(source_row, "client")
            if source_client and source_client.lower() != client.lower():
                raise ValidationError(f"CSV row {number} belongs to {source_client!r}, not {client!r}")
            row = {field: input_value(source_row, field) for field in WORKBOOK_COLUMNS}
            row["client"] = client
            try:
                row["creation_date"] = normalized_date(row["creation_date"])
            except ValidationError as error:
                raise ValidationError(f"CSV row {number}: {error}") from error
            rows.append(row)
    if not rows:
        raise ValidationError("CSV contains no data rows")
    validate_dict_rows(rows, input_path.name)
    clients_dir.mkdir(parents=True, exist_ok=True)
    output_path = clients_dir / f"{client.upper()}.xlsx"
    write_workbook(rows, output_path)
    return output_path


def write_workbook(rows: list[dict[str, str]], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = APPROVED_SHEET
    sheet.append(WORKBOOK_COLUMNS)
    for row in rows:
        sheet.append([row[column] for column in WORKBOOK_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    widths = {"A": 14, "B": 52, "C": 22, "D": 18, "E": 24, "F": 24, "G": 24, "H": 28, "I": 16}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    table = Table(displayName="ApprovedUTMs", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(table)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)


def validate_dict_rows(rows: list[dict[str, str]], source: str) -> None:
    seen: dict[tuple[str, ...], int] = {}
    for index, row in enumerate(rows, start=2):
        missing = [field for field in REQUIRED_COLUMNS if not clean(row.get(field))]
        if missing:
            raise ValidationError(f"{source} row {index}: missing required value(s): {', '.join(missing)}")
        if not valid_http_url(clean(row["destination_url"])):
            raise ValidationError(f"{source} row {index}: invalid destination_url {row['destination_url']!r}")
        bitly = clean(row.get("bitly"))
        if bitly and not valid_http_url(bitly):
            raise ValidationError(f"{source} row {index}: invalid bitly URL {bitly!r}")
        key = tuple(clean(row.get(field)).lower() for field in (
            "client", "destination_url", "source", "medium", "campaign", "term", "content",
        ))
        if key in seen:
            raise ValidationError(f"{source} row {index}: duplicate approved combination (first found on row {seen[key]})")
        seen[key] = index


def read_workbook(path: Path) -> list[ApprovedRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if APPROVED_SHEET not in workbook.sheetnames:
            raise ValidationError(f"{path.name}: missing sheet {APPROVED_SHEET!r}")
        sheet = workbook[APPROVED_SHEET]
        values = sheet.iter_rows(values_only=True)
        try:
            header = tuple(clean(value).lower() for value in next(values))
        except StopIteration as error:
            raise ValidationError(f"{path.name}: workbook is empty") from error
        raw_rows = list(values)
    finally:
        workbook.close()
    if header != WORKBOOK_COLUMNS:
        raise ValidationError(f"{path.name}: expected columns {', '.join(WORKBOOK_COLUMNS)}")
    dict_rows = []
    for excel_row, values_row in enumerate(raw_rows, start=2):
        if not any(clean(value) for value in values_row):
            continue
        row = dict(zip(WORKBOOK_COLUMNS, (clean(value) for value in values_row)))
        try:
            row["creation_date"] = normalized_date(row["creation_date"])
        except ValidationError as error:
            raise ValidationError(f"{path.name} row {excel_row}: {error}") from error
        dict_rows.append(row)
    if not dict_rows:
        raise ValidationError(f"{path.name}: workbook contains no approved rows")
    validate_dict_rows(dict_rows, path.name)
    clients = {row["client"].lower() for row in dict_rows}
    expected_client = path.stem.lower()
    if len(clients) != 1:
        raise ValidationError(f"{path.name}: workbook mixes multiple clients")
    if next(iter(clients)) != expected_client:
        raise ValidationError(f"{path.name}: client values must match workbook name {path.stem!r}")
    return [ApprovedRow(**row, source_file=path.name, excel_row=index) for index, row in enumerate(dict_rows, start=2)]


def load_all_workbooks(clients_dir: Path) -> list[ApprovedRow]:
    paths = sorted(clients_dir.glob("*.xlsx"), key=lambda path: path.name.lower())
    if not paths:
        raise ValidationError(f"no approved .xlsx workbooks found in {clients_dir}")
    rows: list[ApprovedRow] = []
    global_seen: dict[tuple[str, ...], str] = {}
    for path in paths:
        for row in read_workbook(path):
            if row.key in global_seen:
                raise ValidationError(f"{path.name} row {row.excel_row}: duplicates {global_seen[row.key]}")
            global_seen[row.key] = f"{path.name} row {row.excel_row}"
            rows.append(row)
    return rows


def count_entries(rows: Iterable[ApprovedRow], field: str) -> list[dict[str, int | str]]:
    counts: Counter[str] = Counter()
    clients: defaultdict[str, set[str]] = defaultdict(set)
    for row in rows:
        value = row.normalized(getattr(row, field))
        if value:
            counts[value] += 1
            clients[value].add(row.normalized(row.client))
    return [
        {"value": value, "count": count, "unique_clients": len(clients[value])}
        for value, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    ]


def relationship(rows: Iterable[ApprovedRow], left: str, right: str) -> dict[str, list[dict[str, int | str]]]:
    pairs: defaultdict[str, list[tuple[str, str]]] = defaultdict(list)
    for row in rows:
        key = row.normalized(getattr(row, left))
        value = row.normalized(getattr(row, right))
        if key and value:
            pairs[key].append((value, row.normalized(row.client)))
    result = {}
    for key in sorted(pairs):
        counts = Counter(value for value, _client in pairs[key])
        clients = defaultdict(set)
        for value, client in pairs[key]:
            clients[value].add(client)
        result[key] = [
            {"value": value, "count": count, "unique_clients": len(clients[value])}
            for value, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
        ]
    return result


def build_dictionary(rows: list[ApprovedRow]) -> dict[str, object]:
    return {
        "value_counts": {field: count_entries(rows, field) for field in UTM_FIELDS},
        "campaign_to_sources": relationship(rows, "campaign", "source"),
        "campaign_to_mediums": relationship(rows, "campaign", "medium"),
        "source_to_mediums": relationship(rows, "source", "medium"),
        "campaign_to_terms": relationship(rows, "campaign", "term"),
        "campaign_to_contents": relationship(rows, "campaign", "content"),
    }


def csv_text(columns: tuple[str, ...], rows: Iterable[dict[str, object]]) -> str:
    from io import StringIO
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def generated_outputs(rows: list[ApprovedRow]) -> dict[Path, str]:
    dictionary = build_dictionary(rows)
    value_rows = []
    for field in UTM_FIELDS:
        for entry in dictionary["value_counts"][field]:
            value_rows.append({
                "field": field,
                field: entry["value"],
                "count": entry["count"],
                "unique_clients": entry["unique_clients"],
            })
    master_rows = [row.master_dict() for row in sorted(rows, key=lambda item: item.key)]
    return {
        PROJECT_ROOT / "utm_dictionary_output" / "utm_ui_dictionaries.json": json.dumps(dictionary, indent=2, ensure_ascii=False) + "\n",
        PROJECT_ROOT / "utm_dictionary_output" / "utm_value_counts.csv": csv_text(VALUE_COUNT_COLUMNS, value_rows),
        PROJECT_ROOT / "utm_clean_output" / "utm_master_clean.csv": csv_text(MASTER_COLUMNS, master_rows),
    }


def publish(outputs: dict[Path, str]) -> None:
    staged: list[tuple[Path, Path]] = []
    backups: dict[Path, Path | None] = {}
    replaced: list[Path] = []
    try:
        for destination, content in outputs.items():
            destination.parent.mkdir(parents=True, exist_ok=True)
            descriptor, temporary_name = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent)
            with os.fdopen(descriptor, "w", encoding="utf-8", newline="") as handle:
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            staged.append((Path(temporary_name), destination))
        for _temporary, destination in staged:
            if destination.exists():
                descriptor, backup_name = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=".bak", dir=destination.parent)
                os.close(descriptor)
                backup = Path(backup_name)
                shutil.copy2(destination, backup)
                backups[destination] = backup
            else:
                backups[destination] = None
        for temporary, destination in staged:
            os.replace(temporary, destination)
            replaced.append(destination)
    except Exception:
        for destination in reversed(replaced):
            backup = backups.get(destination)
            if backup is None:
                destination.unlink(missing_ok=True)
            else:
                os.replace(backup, destination)
        raise
    finally:
        for temporary, _destination in staged:
            temporary.unlink(missing_ok=True)
        for backup in backups.values():
            if backup is not None:
                backup.unlink(missing_ok=True)


def report(rows: list[ApprovedRow], action: str) -> None:
    clients = Counter(row.client.upper() for row in rows)
    print(f"{action}: {len(rows)} approved UTM rows across {len(clients)} client(s)")
    for client, count in sorted(clients.items()):
        print(f"  {client}: {count}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--clients-dir", type=Path, default=DEFAULT_CLIENTS_DIR, help="approved workbook directory")
    commands = parser.add_subparsers(dest="command", required=True)
    importer = commands.add_parser("import", help="create or replace one approved client workbook from CSV")
    importer.add_argument("--input", type=Path, required=True)
    importer.add_argument("--client", required=True)
    commands.add_parser("validate", help="validate all approved workbooks without writing outputs")
    commands.add_parser("rebuild", help="validate workbooks and replace the app dictionary outputs")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "import":
            output = import_csv(args.input.resolve(), args.client, args.clients_dir.resolve())
            rows = read_workbook(output)
            print(f"Imported {len(rows)} approved rows to {output}")
            return 0
        rows = load_all_workbooks(args.clients_dir.resolve())
        if args.command == "validate":
            report(rows, "Validated")
            return 0
        outputs = generated_outputs(rows)
        publish(outputs)
        report(rows, "Rebuilt")
        for path in outputs:
            print(f"  wrote {path}")
        return 0
    except (OSError, ValidationError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
